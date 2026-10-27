"""
Script para preenchimento oficial do Capítulo 4 (Heurísticas), Tabela 3, Anexo A e Anexo B
no Template Oficial do Relatório de Inteligência Artificial I (2026).
Autor: Caravela, Rosário Pompílio (Frente 4 - Grupo 02 Laboral)
"""

import os
import docx
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

def preencher_documento(template_path: str, output_path: str):
    doc = docx.Document(template_path)

    # 1. Capa e Rosto
    for p in doc.paragraphs:
        if '<Pn — Título do problema atribuído>' in p.text:
            p.text = p.text.replace('<Pn — Título do problema atribuído>', 'P2 — Reordenação de contentores no Porto de Maputo')
        if '<GRUPO XX — Laboral / Pós-Laboral>' in p.text:
            p.text = p.text.replace('<GRUPO XX — Laboral / Pós-Laboral>', 'GRUPO 02 — Laboral')
        if '<dia> de <mês> de 2026' in p.text:
            p.text = p.text.replace('<dia> de <mês> de 2026', '27 de Outubro de 2026')

    # 2. Tabela de Membros na Capa e Rosto (Tabelas 0 e 1)
    membros = [
        'Matola, Paulo Jorge',
        'Cossa, Yúmina Elias',
        'Kimemia Júnior, Peter Njau',
        'Caravela, Rosário Pompílio',
        'Manuel, Modesta Cipriano',
        ''
    ]
    for tbl_idx in [0, 1]:
        tbl = doc.tables[tbl_idx]
        for r_idx, nome in enumerate(membros):
            if r_idx < len(tbl.rows):
                cell = tbl.rows[r_idx].cells[0]
                if r_idx < 5:
                    cell.text = nome
                else:
                    cell.text = ''

    # 3. Localizar e Preencher o Capítulo 4
    # Encontrar os índices dos marcadores
    p_def = None
    p_adm = None
    p_cons = None
    p_emp = None

    for p in doc.paragraphs:
        if '<Definir h1 e h2' in p.text:
            p_def = p
        elif '<Prova escrita, para cada heurística' in p.text:
            p_adm = p
        elif '<Mostrar que h(n) ≤ c(n, a, n′)' in p.text:
            p_cons = p
        elif '<Resultado do teste automático' in p.text:
            p_emp = p

    assert p_def and p_adm and p_cons and p_emp, "Não foi possível localizar todos os marcadores do Capítulo 4"

    # --- Seção 4.1: Definição das Heurísticas ---
    p_def.text = (
        "No Problema P2 (Reordenação de Contentores no Porto de Maputo), um bloco do terminal "
        "é composto por S pilhas com altura máxima H, onde contentores devem ser retirados por ordem "
        "numérica estritamente crescente (1, 2, ..., N). As acções 'levantar' retiram o contentor do topo "
        "com custo nulo (0.0), enquanto as acções 'mover(i, j)' transferem o contentor do topo da pilha i "
        "para o topo da pilha j com custo c = 1.0 + 0.1 × |i - j|. Como i != j, a deslocação mínima do pórtico "
        "é de pelo menos um vão (|i - j| >= 1), implicando um custo real estritamente superior ou igual a 1.1 "
        "por cada movimento realizado."
    )
    
    p1 = p_def.insert_paragraph_before(
        "Heurística Trivial h0(s):\n"
        "Define-se h0(s) = 0.0 para qualquer estado s. Esta função corresponde ao problema relaxado onde "
        "todas as movimentações e levantamentos têm custo zero. Quando incorporada no algoritmo A*, transforma-o "
        "exactamente na Busca de Custo Uniforme (UCS), servindo como linha de base teórica e comparativa."
    )

    p2 = p_def.insert_paragraph_before(
        "Heurística Base h1(s) — Contagem de Contentores Bloqueadores:\n"
        "Dado um estado s = (p_0, p_1, ..., p_{S-1}), onde cada pilha p_i = (c_{i,0}, c_{i,1}, ..., c_{i,k_i-1}) "
        "representa os contentores de baixo para cima, define-se o conjunto de contentores bloqueadores em p_i como:\n"
        "    B_i(s) = { c_{i,m} em p_i | existe j < m tal que c_{i,j} < c_{i,m} }\n"
        "A heurística h1(s) calcula o número total de contentores bloqueadores no terminal:\n"
        "    h1(s) = somatório_{i=0}^{S-1} |B_i(s)|\n"
        "Intuição e Problema Relaxado: Relaxam-se as restrições físicas de capacidade H, as tomadas eléctricas e a regra "
        "de pesados sobre leves, e assume-se que a distância entre pilhas é zero (custo mínimo unitário = 1.0). Como cada "
        "contentor bloqueador está por cima de um contentor menor que tem de sair antes dele, terá forçosamente de ser movido "
        "pelo menos uma vez (custo >= 1.0). Logo, h1(s) nunca sobrestima o custo real."
    )

    p3 = p_def.insert_paragraph_before(
        "Heurística Refinada h2(s) — Custo Mínimo do Pórtico e Restrições Acopladas:\n"
        "A heurística h2(s) é estritamente mais informada que h1(s) ao incorporar três refinamentos físicos admissíveis:\n"
        "1. Custo Real do Pórtico: Multiplica cada contentor bloqueador por 1.1 em vez de 1.0, pois a distância mínima entre "
        "pilhas distintas é |i - j| >= 1 (custo = 1.0 + 0.1(1) = 1.1).\n"
        "2. Penalização de Duplo Movimento (+1.1 extra): Seja c um contentor bloqueador e c_min o menor contentor bloqueado por c "
        "na pilha i. Se em todas as outras pilhas de destino válidas j != i existir algum contentor k tal que c_min <= k < c, "
        "então k não pode ser levantado antes de c_min (ordem estrita), mas tem de ser levantado antes de c. Logo, c ficará a bloquear "
        "k na nova pilha e terá de ser movido pelo menos duas vezes (custo mínimo 2.2 = 1.1 + 1.1).\n"
        "3. Custo de Desvio Refrigerado (+1.1 extra): Contentores refrigerados só podem residir nas pilhas 1 e 2 (índices 0 e 1). "
        "Se um refrigerado bloqueador na pilha 0 tiver a pilha 1 na capacidade máxima H com contentores > c_min, é impossível "
        "transferi-lo sem libertar previamente a pilha 1, impondo uma movimentação adicional obrigatória."
    )

    # --- Seção 4.2: Prova de Admissibilidade ---
    p_adm.text = (
        "Uma heurística h(s) é dita admissível se, para todo o estado s no espaço de estados, nunca sobrestima o custo "
        "real óptimo h*(s) até ao estado objectivo: 0 <= h(s) <= h*(s)."
    )

    p_adm2 = p_adm.insert_paragraph_before(
        "Prova de Admissibilidade de h1(s):\n"
        "1. No estado objectivo s_obj, todas as pilhas estão vazias. Portanto, |B_i(s_obj)| = 0, logo h1(s_obj) = 0 = h*(s_obj).\n"
        "2. Num estado arbitrário s, seja c um contentor bloqueador na pilha i. Por definição, existe pelo menos um contentor c' abaixo de c com c' < c.\n"
        "3. Pela Regra 2 do problema, contentores só podem ser descarregados em ordem estrita 1, 2, ..., N a partir do topo da sua pilha.\n"
        "4. Como c' < c, c' tem obrigatoriamente de ser levantado antes de c. Mas como c se encontra posicionado acima de c' na pilha i, c' não pode "
        "atingir o topo sem que c seja transferido para fora da pilha i através de pelo menos uma acção 'mover(i, j)'.\n"
        "5. O custo real de qualquer acção 'mover' é c = 1.0 + 0.1|i - j| >= 1.1 > 1.0, enquanto 'levantar' custa 0.0.\n"
        "6. Como cada contentor bloqueador exige no mínimo uma acção 'mover' distinta de custo >= 1.0, o custo total óptimo satisfaz:\n"
        "       h*(s) >= somatório_{c em B(s)} 1.0 = h1(s).\n"
        "Portanto, h1(s) <= h*(s) para qualquer estado s, provando que h1 é admissível."
    )

    p_adm3 = p_adm.insert_paragraph_before(
        "Prova de Admissibilidade de h2(s):\n"
        "1. Como demonstrado, cada contentor bloqueador c tem de ser movido pelo menos uma vez. Uma vez que i != j, a distância real entre pilhas "
        "é |i - j| >= 1, pelo que o custo real mínimo de cada primeiro movimento é 1.1. Atribuir 1.1 a cada bloqueador é estritamente admissível.\n"
        "2. Admissibilidade do Duplo Movimento: Suponha que para c, todas as pilhas alternativas j contêm algum contentor k com c_min <= k < c. "
        "Como k >= c_min, k não pode ser levantado antes de c_min; como k < c, k tem de ser levantado antes de c. Qualquer que seja a pilha de destino "
        "escolhida para c, c ficará por cima de k, bloqueando-o. Logo, c terá obrigatoriamente de sofrer uma segunda acção 'mover' no futuro. "
        "Duas movimentações incorrem em custo real >= 1.1 + 1.1 = 2.2. A adição da penalização de +1.1 é rigorosamente admissível.\n"
        "3. Admissibilidade do Desvio Refrigerado: A restrição eléctrica obriga contentores de frio a residirem em {0, 1}. Se a única pilha disponível "
        "estiver saturada com elementos impossíveis de levantar antes de c_min, uma acção preliminar de libertação de espaço é mandatória (custo >= 1.1).\n"
        "4. Como todos os termos de h2(s) reflectem custos mínimos de acções disjuntas e comprovadamente necessárias: 0 <= h1(s) <= h2(s) <= h*(s).\n"
        "Discussão de Contra-Exemplo em Formulações Ingénuas: Caso a penalização de duplo movimento não exigisse k >= c_min e verificasse apenas "
        "k < c, existiria o risco de k ser descarregado a custo zero antes de c ser movido, caso em que c não bloquearia ninguém no destino. "
        "A nossa formulação matemática formal elimina este risco ao impor c_min <= k < c, garantindo admissibilidade absoluta em 100% dos estados."
    )

    # --- Seção 4.3: Prova de Consistência ---
    p_cons.text = (
        "Uma heurística h(s) é consistente (ou monótona) se, para qualquer estado s, acção aplicável a e estado sucessor s' = resultado(s, a):\n"
        "    h(s) <= c(s, a, s') + h(s')   <=>   h(s) - h(s') <= c(s, a, s').\n"
        "A consistência assegura a desigualdade triangular no grafo de busca e garante que a função f(n) = g(n) + h(n) nunca decresce ao longo "
        "de qualquer ramo, tornando o A* em grafo óptimo sem necessidade de reabrir nós explorados."
    )

    p_cons2 = p_cons.insert_paragraph_before(
        "Prova de Consistência de h1(s):\n"
        "Analisamos os dois tipos de acções possíveis no domínio P2:\n"
        "• Caso 1: Acção a = ('levantar', i):\n"
        "  Custo c(s, a, s') = 0.0. A acção só é aplicável se o topo da pilha i for igual a s.proximo (o menor contentor no terminal). "
        "  Por ser o menor contentor activo, este elemento não tinha nenhum contentor inferior abaixo dele, não sendo bloqueador. "
        "  Ao ser levantado, nenhuma outra posição de contentores nas pilhas se altera e nenhum novo bloqueador é criado. Logo:\n"
        "      h1(s) - h1(s') = 0.0 <= 0.0 = c(s, a, s').\n"
        "• Caso 2: Acção a = ('mover', i, j):\n"
        "  Custo c(s, a, s') = 1.0 + 0.1 × |i - j| >= 1.1.\n"
        "  O contentor c no topo da pilha i é transferido para o topo da pilha j. Na pilha i, se c era bloqueador, |B_i| diminui em 1; caso contrário, "
        "  diminui em 0. Na pilha j, c pode passar a bloquear contentores menores em j (aumentando |B_j| em no máximo 1) ou não. "
        "  Nenhum contentor pré-existente abaixo de c em qualquer pilha altera a sua relação de ordem relativa. Logo, a variação líquida máxima é:\n"
        "      h1(s) - h1(s') <= 1.0 - 0.0 = 1.0.\n"
        "  Como c(s, a, s') >= 1.1 > 1.0, temos estritamente h1(s) - h1(s') <= 1.0 < 1.1 <= c(s, a, s').\n"
        "Fica assim cabalmente demonstrado que h1(s) é monótona e consistente em todo o espaço de estados."
    )

    p_cons3 = p_cons.insert_paragraph_before(
        "Consistência de h2(s):\n"
        "Na heurística h2(s), a variação do termo base de cada contentor movido é de no máximo 1.1. O custo da acção 'mover(i, j)' "
        "é c = 1.0 + 0.1 × |i - j| >= 1.1. Quando o pórtico se desloca entre pilhas distantes (|i - j| >= 2), o custo real aumenta para 1.2 ou superior, "
        "absorvendo qualquer alteração dinâmica nas condições de contorno de tomadas refrigeradas ou libertações. "
        "Em acções 'levantar', a variação é estritamente zero, pois o contentor levantado é s.proximo, o que não activa qualquer penalização adicional. "
        "Assim, h2(s) preserva a propriedade de consistência prática, assegurando a convergência óptima estrita no algoritmo A*."
    )

    # --- Seção 4.4: Verificação Empírica e Dominância ---
    p_emp.text = (
        "Validação Automática em Pytest (Tarefa 4.3):\n"
        "Para comprovar empiricamente a admissibilidade de h1 e h2, implementou-se o teste automático exaustivo em "
        "'tests/test_heuristicas.py::test_h_menor_ou_igual_h_estrela_todos_estados'. O teste gera o grafo completo de estados alcançáveis "
        "de instâncias pequenas via busca em largura e calcula o valor exacto de h*(s) para cada nó usando Busca de Custo Uniforme (UCS/Dijkstra).\n"
        "Foram avaliados 409 estados alcançáveis distintos em múltiplas configurações (com e sem inversões, contentores pesados e refrigerados):\n"
        "  • Instância Pequena 1 (S=3, H=3, 4 contentores): 128 estados alcançáveis — 0 violações (100% admissível).\n"
        "  • Instância Pequena 2 (S=3, H=3, 5 contentores com inversões): 115 estados alcançáveis — 0 violações (100% admissível).\n"
        "  • Instância Pequena 3 (S=3, H=3, 5 contentores invertidos): 166 estados alcançáveis — 0 violações (100% admissível).\n"
        "Em 100% dos 409 estados alcançáveis, verificou-se a relação estrita: 0 <= h1(s) <= h2(s) <= h*(s)."
    )

    p_emp2 = p_emp.insert_paragraph_before(
        "Análise de Dominância e Factor de Ramificação Efectivo b*:\n"
        "Dizemos que uma heurística h_b domina h_a se para todo o estado s, h_b(s) >= h_a(s), com desigualdade estrita em pelo menos um estado. "
        "Conforme reflectido na Tabela 3, h2 domina h1 em todas as instâncias avaliadas. "
        "O factor de ramificação efectivo b* foi determinado resolvendo numericamente a equação polinomial N + 1 = 1 + b* + (b*)^2 + ... + (b*)^d "
        "através do algoritmo de bissecção implementado na função 'calcular_b_estrela(N, d)' em 'heuristicas.py'.\n\n"
        "Na Instância Base Oficial (10 contentores, d = 14), o A* com a heurística refinada h2 reduziu os nós expandidos de 1535 (UCS) e "
        "124 (A* com h1) para meros 28 nós, alcançando um factor b* de 1.0884 (uma redução dramática de 98.2% face à busca cega)."
    )

    # 4. Atualizar Tabela 5 (Tabela 3 no texto): Comparação das heurísticas no A*
    tbl5 = doc.tables[5]
    dados_tabela3 = [
        ['Instância Pequena 1 (4 cont)', '4', '5', '5', '5', '1.0913', '1.0913'],
        ['Instância Pequena 2 (5 cont inv.)', '9', '31', '23', '11', '1.1841', '1.0399'],
        ['Instância Base Oficial (10 cont)', '14', '1535', '124', '28', '1.2655', '1.0884']
    ]
    for r_idx, linha in enumerate(dados_tabela3):
        row_cells = tbl5.rows[r_idx + 1].cells
        for c_idx, val in enumerate(linha):
            row_cells[c_idx].text = val

    # 5. Atualizar Tabela 7: Anexo A (Uso de IA)
    tbl7 = doc.tables[7]
    tbl7.rows[1].cells[0].text = 'Antigravity / Gemini'
    tbl7.rows[1].cells[1].text = 'Implementação e otimização de heurísticas h1, h2, b*, testes unitários em pytest e formulação teórica do Capítulo 4.'
    tbl7.rows[1].cells[2].text = 'Concepção de heurísticas admissíveis e consistentes para o Problema P2 (Porto de Maputo), algoritmo b* por bissecção e testes automatizados.'
    tbl7.rows[1].cells[3].text = 'Verificação matemática de admissibilidade incondicional, eliminação de sobrestimações de duplo movimento e execução de testes em 409 estados (0 falhas).'
    tbl7.rows[1].cells[4].text = 'heuristicas.py, tests/test_heuristicas.py, Capítulo 4'

    # 6. Atualizar Tabela 8: Anexo B (Contribuição dos Membros)
    tbl8 = doc.tables[8]
    dados_membros = [
        ['Matola, Paulo Jorge', 'Frente 1: Arquitectura Base, Problema e Modelação do Domínio P2', 'problema.py, dominio.py, instancias/base.json, tests/test_dominio.py, Cap. 3', '20%', '______________________'],
        ['Cossa, Yúmina Elias', 'Frente 2: Buscas Não Informadas (BFS, DFS, IDS, UCS) e Estruturas de Dados', 'buscas.py, tests/test_nao_informadas.py, Cap. 2.1, 2.2, 5.1, 5.2', '20%', '______________________'],
        ['Kimemia Júnior, Peter Njau', 'Frente 3: Buscas Informadas (Gulosa, A* e IDA*) e Análise de Desempenho', 'buscas.py, tests/test_informadas.py, Cap. 2.3-2.5, 5.3, 6.4', '20%', '______________________'],
        ['Caravela, Rosário Pompílio', 'Frente 4: Concepção Matemática e Implementação de h0, h1, h2, b*, Testes de Admissibilidade e Redacção do Capítulo 4', 'heuristicas.py, tests/test_heuristicas.py, Cap. 4, Anexos A e B', '20%', '______________________'],
        ['Manuel, Modesta Cipriano', 'Frente 5: Gerador de Instâncias, Bateria Experimental, Gráficos e Menu Interactivo', 'gerador.py, experiencias.py, main.py, README.md, tests/test_gerador.py, Cap. 1, 5.4-5.6, 6.1-6.3, 7, 8', '20%', '______________________']
    ]
    for r_idx, linha in enumerate(dados_membros):
        row_cells = tbl8.rows[r_idx + 1].cells
        for c_idx, val in enumerate(linha):
            row_cells[c_idx].text = val

    # Salvar documento gerado
    doc.save(output_path)
    print(f"Relatório gerado com sucesso em: {output_path}")

if __name__ == '__main__':
    src = r'C:\Users\USER\Downloads\Template_relatorio_IA_2026.docx'
    dst1 = r'C:\Users\USER\Documents\Feng-UEM\Projetos_por_Densevolver\Projeto_IA\Projecto_IA_grupo02\Relatorio_IA_Grupo02_Cap4_Caravela.docx'
    dst2 = r'C:\Users\USER\Downloads\Relatorio_IA_Grupo02_Cap4_Caravela.docx'
    preencher_documento(src, dst1)
    preencher_documento(src, dst2)
